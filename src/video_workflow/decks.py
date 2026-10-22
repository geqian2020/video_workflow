import os
import shutil
import subprocess
import zipfile
from pathlib import Path

from .models import Slide

MAX_PAGES = 60


def check_pptx(path: Path):
    try:
        with zipfile.ZipFile(path) as archive:
            files = archive.infolist()
            if len(files) > 10000 or sum(f.file_size for f in files) > 512 * 1024 * 1024:
                raise ValueError("PPT 解压后过大，请拆分后导入")
            if "ppt/presentation.xml" not in archive.namelist():
                raise ValueError("文件不是有效的 PPTX")
    except zipfile.BadZipFile as exc:
        raise ValueError("PPTX 已损坏或格式不正确") from exc


def render_pdf(path: Path, output: Path) -> tuple[list[Path], list[str]]:
    import pypdfium2 as pdfium

    output.mkdir(parents=True, exist_ok=True)
    files, texts = [], []
    with pdfium.PdfDocument(str(path)) as pdf:
        if not 1 <= len(pdf) <= MAX_PAGES:
            raise ValueError(f"请使用 1–{MAX_PAGES} 页的 PDF")
        for index in range(len(pdf)):
            page = pdf[index]
            try:
                scale = min(1920 / max(page.get_size()), 3)
                bitmap = page.render(scale=scale)
                try:
                    file = output / f"page-{index + 1:03}.png"
                    bitmap.to_pil().save(file)
                finally:
                    bitmap.close()
                textpage = page.get_textpage()
                try:
                    texts.append(textpage.get_text_range())
                finally:
                    textpage.close()
                files.append(file)
            finally:
                page.close()
    return files, texts


def render_powerpoint(path: Path, output: Path) -> list[Path]:
    """Read-only Office export on Windows; isolated LibreOffice profile otherwise."""
    output.mkdir(parents=True, exist_ok=True)
    if os.name == "nt":
        from .windows import run_powershell

        try:
            process = run_powershell(
                "export_ppt.ps1",
                {
                    "WORKFLOW_INPUT_DECK": str(path.resolve()),
                    "WORKFLOW_OUTPUT_DIR": str(output.resolve()),
                },
                timeout=180,
            )
            files = sorted(output.glob("page-*.png"))
            if process.returncode == 0 and files:
                return files
        except subprocess.TimeoutExpired:
            pass
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise ValueError(
            "未能调用 PowerPoint。请安装 PowerPoint/LibreOffice，或把 PPT 导出 PDF 后上传。"
        )
    profile = output / "lo-profile"
    process = subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={profile.resolve().as_uri()}",
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output),
            str(path.resolve()),
        ],
        capture_output=True,
        timeout=180,
    )
    pdf = output / f"{path.stem}.pdf"
    if process.returncode or not pdf.exists():
        raise ValueError("PPT 转换失败，请先在 PowerPoint 中导出 PDF")
    return render_pdf(pdf, output)[0]


def import_deck(path: Path, root: Path) -> tuple[list[Slide], list[str]]:
    output = root / "slides"
    warnings = []
    if path.suffix.lower() == ".pptx":
        check_pptx(path)
        # Reuse SlideStream's title/content/notes extraction, but preserve every
        # original page, including completely blank pages it normally omits.
        from pptx import Presentation
        from slide_stream.powerpoint import parse_powerpoint

        presentation = Presentation(str(path))
        if not 1 <= len(presentation.slides) <= MAX_PAGES:
            raise ValueError(f"请使用 1–{MAX_PAGES} 页的 PPTX")
        parsed = iter(parse_powerpoint(path))
        metadata = []
        for page in presentation.slides:
            meaningful = any(shape.has_text_frame and shape.text.strip() for shape in page.shapes)
            for shape in page.shapes:
                if shape.shape_type == 13:
                    try:
                        meaningful = meaningful or bool(shape.image.blob)
                    except Exception:
                        pass
            if page.has_notes_slide:
                notes_frame = page.notes_slide.notes_text_frame
                if notes_frame is not None:
                    note = notes_frame.text.strip()
                    meaningful = meaningful or bool(note and note != "Click to add notes")
                else:
                    meaningful = meaningful or any(
                        shape.has_text_frame
                        and shape.text.strip()
                        and shape.text.strip() != "Click to add notes"
                        for shape in page.notes_slide.shapes
                    )
            metadata.append(next(parsed, {}) if meaningful else {})
        files = render_powerpoint(path, output)
        if len(files) != len(metadata):
            raise ValueError("PPT 页面导出数量不一致，请改用 PDF，避免配音与页面错位")
        warnings.append("保留静态页面画面；原 PPT 的动画和内嵌视频不会自动播放。")
    elif path.suffix.lower() == ".pdf":
        files, texts = render_pdf(path, output)
        metadata = [
            {
                "title": text.strip().splitlines()[0][:120] if text.strip() else f"第 {i + 1} 页",
                "content": text.strip().splitlines(),
                "notes": "",
            }
            for i, text in enumerate(texts)
        ]
    else:
        raise ValueError("只接受 PPTX 或 PDF")
    slides = []
    for index, (file, data) in enumerate(zip(files, metadata, strict=True)):
        title = str(data.get("title") or f"第 {index + 1} 页")
        content = list(data.get("content") or [])
        notes = str(data.get("notes") or "").strip()
        slides.append(
            Slide(
                title=title,
                content=content,
                narration=(notes or "\n".join(content))[:8000],
                image=file.relative_to(root).as_posix(),
                source_page=index + 1,
            )
        )
    if any(not slide.narration.strip() for slide in slides):
        warnings.append("有页面没有可提取的文字，请填写讲稿，或接入视觉模型后人工核对。")
    warnings.append("初始讲稿来自备注或页面文字，请编辑核对后配音。")
    return slides, warnings
