import io
import time
import wave

import numpy as np
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation

from video_workflow.app import create_app
from video_workflow.models import Avatar, Project, RenderOptions, Slide
from video_workflow.speech import Cue, sentences, srt_timestamp, synthesize, write_srt
from video_workflow.storage import Store
from video_workflow.video import envelope


def tiny_pdf():
    # A deterministic blank PDF fixture, not a user-facing presentation.
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 960 540] /Resources << >> >>",
    ]
    data = b"%PDF-1.4\n"
    offsets = []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{index} 0 obj\n".encode() + obj + b"\nendobj\n"
    start = len(data)
    data += b"xref\n0 4\n0000000000 65535 f \n"
    for offset in offsets:
        data += f"{offset:010d} 00000 n \n".encode()
    return data + f"trailer\n<< /Size 4 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()


def wait_ready(client, project_id):
    for _ in range(300):
        result = client.get(f"/api/projects/{project_id}").json()
        if result["status"] not in {"importing", "generating", "rendering", "previewing"}:
            return result
        time.sleep(0.01)
    pytest.fail("background task did not finish")


@pytest.fixture
def client(tmp_path):
    with TestClient(create_app(tmp_path)) as value:
        yield value


def imported(client):
    response = client.post(
        "/api/projects", files={"file": ("原稿.pdf", tiny_pdf(), "application/pdf")}
    )
    assert response.status_code == 202
    result = wait_ready(client, response.json()["id"])
    assert result["status"] == "ready", result
    return result


def test_import_pdf_edit_persist_and_restrict_assets(client):
    project = imported(client)
    assert len(project["slides"]) == 1
    slide = project["slides"][0]
    prefix = f"/api/projects/{project['id']}"
    assert client.get(prefix + "/files/" + slide["image"]).headers["content-type"] == "image/png"
    response = client.patch(
        prefix + f"/slides/{slide['id']}", json={"title": "整理后的标题", "narration": "逐页讲解。"}
    )
    assert response.status_code == 200
    assert response.json()["version"] == project["version"] + 1
    assert client.get(prefix).json()["slides"][0]["narration"] == "逐页讲解。"
    assert client.get(prefix + "/files/project.json").status_code == 404
    assert client.get(prefix + "/files/source/source.pdf").status_code == 404


def test_invalid_file_and_missing_project(client):
    assert client.post("/api/projects", files={"file": ("bad.exe", b"bad")}).status_code == 400
    response = client.post("/api/projects", files={"file": ("bad.pptx", b"bad")})
    result = wait_ready(client, response.json()["id"])
    assert result["status"] == "failed"
    assert "PPTX" in result["error"]
    assert client.get("/api/projects/../bad").status_code == 404
    assert client.get("/api/projects/" + "a" * 32).status_code == 404


def test_no_external_origin_or_host_mutations(client):
    response = client.post("/api/projects", headers={"Origin": "https://other.example"})
    assert response.status_code == 403
    assert client.get("/api/projects", headers={"host": "other.example"}).status_code == 400


def test_generation_requires_config(client, monkeypatch):
    monkeypatch.delenv("PRESENTON_URL", raising=False)
    response = client.post("/api/generate", json={"content": "这是一份有足够长度的中文课程文案。"})
    assert response.status_code == 503
    assert "PRESENTON_URL" in response.json()["detail"]


def test_avatar_rejects_invalid_image_and_path_edit(client):
    project = imported(client)
    prefix = f"/api/projects/{project['id']}"
    assert (
        client.post(prefix + "/avatar", files={"file": ("image.png", b"not an image")}).status_code
        == 400
    )
    image = Image.new("RGBA", (200, 300), (0, 0, 0, 0))
    image.putpixel((100, 100), (20, 80, 90, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    response = client.post(prefix + "/avatar", files={"file": ("image.png", buffer.getvalue())})
    assert response.status_code == 200
    body = response.json()["avatar"]
    original = body["image"]
    body["image"] = "../../private.json"
    response = client.patch(prefix + "/avatar", json=body)
    assert response.json()["avatar"]["image"] == original
    assert response.json()["version"] == project["version"] + 1
    assert client.post(prefix + "/render", json={"slide_ids": ["unknown"]}).status_code == 202
    assert wait_ready(client, project["id"])["status"] == "failed"


def test_render_without_character_is_explained(client):
    project = imported(client)
    response = client.post(f"/api/projects/{project['id']}/render", json={})
    assert response.status_code == 400
    assert "角色" in response.json()["detail"]


def test_store_recovery_and_stale_output(tmp_path):
    store = Store(tmp_path)
    original = store.save(
        Project(
            status="rendering",
            version=3,
            output_version=2,
            outputs={"video": "output/previous.mp4"},
        )
    )
    store.recover()
    current = store.load(original.id)
    assert current.status == "failed"
    assert current.outputs == original.outputs
    assert store.manifest(current)["output_stale"] is True
    with pytest.raises(FileNotFoundError):
        store.directory("../../outside")


def test_pptx_blank_pages_keep_original_mapping(tmp_path, monkeypatch):
    from video_workflow.decks import import_deck

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    titled = prs.slides.add_slide(prs.slide_layouts[0])
    titled.shapes.title.text = "第二页"
    titled.notes_slide.notes_text_frame.text = "第二页的讲稿"
    prs.slides.add_slide(prs.slide_layouts[6])
    path = tmp_path / "source.pptx"
    prs.save(path)

    def fake_export(path, output):
        output.mkdir()
        files = []
        for i in range(3):
            image = output / f"page-{i + 1:03}.png"
            Image.new("RGB", (64, 36)).save(image)
            files.append(image)
        return files

    monkeypatch.setattr("video_workflow.decks.render_powerpoint", fake_export)
    slides, _ = import_deck(path, tmp_path)
    assert len(slides) == 3
    assert slides[0].narration == slides[2].narration == ""
    assert slides[1].title == "第二页"
    assert slides[1].narration == "第二页的讲稿"
    assert [s.source_page for s in slides] == [1, 2, 3]


def test_speech_cache_uses_text_and_engine_and_real_duration(tmp_path, monkeypatch):
    called = []

    def fake_voice(text, path):
        called.append(text)
        with wave.open(str(path), "wb") as out:
            out.setnchannels(1)
            out.setsampwidth(2)
            out.setframerate(16000)
            out.writeframes(np.zeros(8000, dtype="<i2").tobytes())

    monkeypatch.setattr("video_workflow.speech._sapi", fake_voice)
    audio, cues = synthesize("第一句。第二句！", tmp_path, "sapi")
    assert [c.text for c in cues] == ["第一句。", "第二句！"]
    assert cues[-1].end == 1.0
    assert audio.exists()
    again, _ = synthesize("第一句。第二句！", tmp_path, "sapi")
    assert again == audio and len(called) == 2
    changed, _ = synthesize("修改了。", tmp_path, "sapi")
    assert changed != audio and len(called) == 3


def test_silence_closes_mouth(tmp_path):
    audio = tmp_path / "silence.wav"
    with wave.open(str(audio), "wb") as out:
        out.setnchannels(1)
        out.setsampwidth(2)
        out.setframerate(16000)
        out.writeframes(np.zeros(16000, dtype="<i2").tobytes())
    assert np.all(envelope(audio, 24) == 0)


def test_srt_and_sentence_content(tmp_path):
    assert "".join(sentences("第一页。\n第二页！")) == "第一页。第二页！"
    assert srt_timestamp(59.9996) == "00:01:00,000"
    path = tmp_path / "captions.srt"
    write_srt([Cue(0, 2.1, "测试字幕。")], path)
    assert "00:00:02,100" in path.read_text(encoding="utf-8-sig")


def test_mouth_region_validation():
    with pytest.raises(ValueError):
        Avatar(mouth=(0.01, 0.5, 0.4, 0.1))


def test_download_does_not_forward_credentials_to_other_origin(tmp_path):
    from video_workflow.generation import download_export

    with pytest.raises(ValueError, match="下载地址"):
        download_export(
            "http://localhost:5001",
            {"Authorization": "Bearer not-real"},
            {"path": "https://other.example/steal"},
            tmp_path / "deck.pptx",
        )


def test_renderer_handles_empty_narration_before_starting_encoding(tmp_path):
    from video_workflow.video import render_project

    project = Project(avatar=Avatar(), slides=[Slide(image="slides/a.png", source_page=1)])
    with pytest.raises(ValueError, match="空讲稿"):
        render_project(project, tmp_path, RenderOptions(), lambda _: None)


def test_script_change_invalidates_only_its_voice_preview(client):
    project = imported(client)
    store = client.app.state.store
    slide = project["slides"][0]
    current = store.load(project["id"])
    current.outputs["preview_audio"] = "audio/old.wav"
    current.preview_audio_slide = slide["id"]
    store.save(current)
    response = client.patch(
        f"/api/projects/{current.id}/slides/{slide['id']}",
        json={"title": slide["title"], "narration": "这是修改后的讲稿。"},
    )
    assert "preview_audio" not in response.json()["outputs"]
    assert response.json()["preview_audio_slide"] is None


def test_page_preview_preserves_existing_full_video(client, monkeypatch):
    project = imported(client)
    store = client.app.state.store
    current = store.load(project["id"])
    current.avatar = Avatar()
    current.outputs["video"] = "output/full.mp4"
    current.output_version = current.version
    store.save(current)
    monkeypatch.setattr(
        "video_workflow.video.render_project",
        lambda *args: {"video": "output/preview.mp4", "subtitles": "output/preview.srt"},
    )
    response = client.post(
        f"/api/projects/{current.id}/render", json={"slide_ids": [current.slides[0].id]}
    )
    result = wait_ready(client, response.json()["id"])
    assert result["outputs"]["video"] == "output/full.mp4"
    assert result["outputs"]["preview_video"] == "output/preview.mp4"
    assert result["output_stale"] is False


def test_long_display_captions_preserve_all_characters():
    from PIL import ImageFont

    from video_workflow.video import display_cues

    source = "字幕需要完整保留，不能因为放不下就丢掉后半段内容。" * 4
    cues = display_cues([Cue(0, 12, source)], 180, ImageFont.load_default(size=20))
    assert "".join(c.text.replace("\n", "") for c in cues) == source
    assert cues[0].start == 0 and cues[-1].end == 12
    assert all(len(c.text.splitlines()) <= 2 for c in cues)
